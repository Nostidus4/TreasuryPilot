"""Export the nine-part English TreasuryPilot report to an editable DOCX.

Install dependencies with: pip install -r docs/docx_requirements.txt
Then run: python docs/export_product_report_docx.py
"""
from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from markdown_it import MarkdownIt

REPO = Path(__file__).resolve().parents[1]
DOCS = REPO / "docs"
SOURCE = DOCS / "PRODUCT_REPORT_EN.md"
OUTPUT = DOCS / "TreasuryPilot_Product_Technical_Report_EN.docx"
SHA = "50b1117"
GITHUB = "https://github.com/Nostidus4/TreasuryPilot"
NAVY = RGBColor(17, 45, 86)
BLUE = RGBColor(32, 88, 190)
PINK = RGBColor(228, 60, 145)
MUTED = RGBColor(73, 100, 128)
PALE = "EAF2FA"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border_bottom(paragraph, color: str = "B8C9E0") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "5"), ("color", color)):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    p_pr.append(borders)


def page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def link_url(raw: str) -> str:
    if raw.startswith(("https://", "http://", "mailto:")):
        return raw
    candidate = (DOCS / raw).resolve()
    if not candidate.is_relative_to(REPO) or not candidate.exists():
        raise ValueError(f"Broken report source link: {raw}")
    relative = candidate.relative_to(REPO).as_posix()
    kind = "tree" if candidate.is_dir() else "blob"
    if relative in {"docs/TECHNICAL_REPORT_EN.md", "docs/TECHNICAL_REPORT_VI.md", "docs/TECHNICAL_REPORT_ZH.md"}:
        return f"{GITHUB}/{kind}/main/{relative}"
    return f"{GITHUB}/{kind}/{SHA}/{relative}"


def hyperlink(paragraph, label: str, raw_url: str) -> None:
    url = link_url(raw_url)
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    outer = OxmlElement("w:hyperlink")
    outer.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2058BE")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    outer.append(run)
    paragraph._p.append(outer)


def render_inline(paragraph, children) -> None:
    bold = italic = False
    active_url = None
    for token in children or []:
        kind = token.type
        if kind == "strong_open": bold = True
        elif kind == "strong_close": bold = False
        elif kind == "em_open": italic = True
        elif kind == "em_close": italic = False
        elif kind == "link_open": active_url = token.attrGet("href")
        elif kind == "link_close": active_url = None
        elif kind in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif kind in {"text", "code_inline"}:
            if active_url:
                hyperlink(paragraph, token.content, active_url)
            else:
                run = paragraph.add_run(token.content)
                run.bold = bold
                run.italic = italic
                if kind == "code_inline":
                    run.font.name = "Menlo"
                    run.font.size = Pt(9)
                    run.font.color.rgb = BLUE


def text_of(children) -> str:
    return "".join(token.content for token in children or [] if token.type in {"text", "code_inline"})


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.16
    for name, size, before, after, color in [
        ("Title", 27, 0, 18, NAVY),
        ("Heading 1", 19, 20, 10, BLUE),
        ("Heading 2", 12.5, 14, 6, NAVY),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
    if "Source Note" not in doc.styles:
        style = doc.styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.size = Pt(8.5)
        style.font.color.rgb = MUTED
        style.paragraph_format.space_after = Pt(7)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(map(len, rows))
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, cells in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(cells):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            if row_index == 0:
                shade(cell, "DDECF7")
            elif row_index % 2 == 0:
                shade(cell, "F5F8FC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.4 if width >= 4 else 9)
                    run.font.color.rgb = NAVY
                    if row_index == 0:
                        run.bold = True
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(tokens, start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    row: list[str] | None = None
    i = start + 1
    while i < len(tokens) and tokens[i].type != "table_close":
        token = tokens[i]
        if token.type == "tr_open": row = []
        elif token.type == "tr_close" and row is not None:
            rows.append(row)
            row = None
        elif token.type == "inline" and row is not None:
            row.append(text_of(token.children))
        i += 1
    return rows, i + 1


def render_report(doc: Document) -> None:
    md = MarkdownIt("default").enable("table")
    tokens = md.parse(SOURCE.read_text(encoding="utf-8"))
    started = False
    section_count = 0
    list_type = None
    in_blockquote = False
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type == "heading_open" and token.tag == "h2":
            started = True
            section_count += 1
            heading = doc.add_heading(text_of(tokens[i+1].children), level=1)
            heading.paragraph_format.page_break_before = True
            i += 3
            continue
        if not started:
            i += 1
            continue
        if token.type == "heading_open" and token.tag == "h3":
            doc.add_heading(text_of(tokens[i+1].children), level=2)
            i += 3
            continue
        if token.type == "table_open":
            rows, i = parse_table(tokens, i)
            add_table(doc, rows)
            continue
        if token.type in {"bullet_list_open", "ordered_list_open"}:
            list_type = "List Bullet" if token.type == "bullet_list_open" else "List Number"
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            list_type = None
        elif token.type == "blockquote_open":
            in_blockquote = True
        elif token.type == "blockquote_close":
            in_blockquote = False
        elif token.type == "paragraph_open" and i + 1 < len(tokens) and tokens[i+1].type == "inline":
            inline = tokens[i+1]
            source_note = inline.content.startswith("Sources:") or inline.content.startswith("Source:")
            paragraph = doc.add_paragraph(style="Source Note" if source_note else list_type)
            if in_blockquote:
                paragraph.paragraph_format.left_indent = Inches(.2)
            render_inline(paragraph, inline.children)
            i += 3
            continue
        elif token.type == "fence":
            paragraph = doc.add_paragraph()
            paragraph.style = "Source Note"
            paragraph.add_run(token.content)
        i += 1
    if section_count != 9:
        raise ValueError(f"Expected nine numbered sections, got {section_count}")


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(.74)
    section.bottom_margin = Inches(.68)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)
    section.header_distance = Inches(.34)
    section.footer_distance = Inches(.35)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.text = "TREASURYPILOT   /   PRODUCT & TECHNICAL REPORT"
    header.style = doc.styles["Source Note"]
    border_bottom(header)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = doc.styles["Source Note"]
    page_number(footer)

    brand = doc.add_paragraph()
    brand.add_run("TREASURYPILOT").font.color.rgb = PINK
    brand.runs[0].bold = True
    brand.runs[0].font.size = Pt(15)
    brand.paragraph_format.space_before = Pt(58)
    title = doc.add_paragraph(style="Title")
    title.add_run("Product & Technical Report")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Invoice-to-stablecoin payment with human approval, AI assistance and constrained execution").italic = True
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = MUTED
    doc.add_paragraph()
    info = doc.add_table(rows=3, cols=2)
    info.style = "Table Grid"
    for row, (label,value) in enumerate([
        ("Prepared", "16 October 2026"),
        ("Code baseline", SHA),
        ("Evidence scope", "Repository code, dated local/fork results and Sepolia manifest"),
    ]):
        info.cell(row,0).text=label
        info.cell(row,1).text=value
        shade(info.cell(row,0), "DDECF7")
    doc.add_paragraph()
    summary = doc.add_paragraph()
    summary.add_run("EXECUTIVE SUMMARY\n").bold = True
    summary.add_run(
        "TreasuryPilot coordinates stablecoin invoice payments across an organization's Safe wallets. "
        "A human confirms and signs the obligation; AI extracts and explains; a deterministic engine plans "
        "liquidity and timing; policy, an isolated signer and an on-chain module constrain execution. "
        "The repository contains substantial implementations and local evidence. Public Sepolia end-to-end "
        "results and independent security review remain separate verification steps."
    )
    summary.paragraph_format.left_indent = Inches(.18)
    summary.paragraph_format.right_indent = Inches(.18)
    summary.paragraph_format.space_before = Pt(28)
    border_bottom(summary, "E43C91")
    note = doc.add_paragraph(style="Source Note")
    note.add_run("Evidence standard: code presence, local/fork tests, deployment configuration and public-chain results are reported separately. ")
    hyperlink(note, "Open project document", "https://docs.google.com/document/d/1nkSIU5c-zouBkKooWjEFshxw61pH0CRwugYDy5AFJBo/edit")

    contents = doc.add_heading("Contents", level=1)
    contents.paragraph_format.page_break_before = True
    for entry in [
        "1. Target Users & Current Journey",
        "2. Pain Points & Solution Directions",
        "3. MVP Scope & Current Coverage",
        "4. Introducing TreasuryPilot",
        "5. How TreasuryPilot Works",
        "6. High-Level System Architecture",
        "7. AI & Core Logic: Why Both AI and an Engine?",
        "8. Security and Permission Design",
        "9. Results & Closing",
    ]:
        doc.add_paragraph(entry)
    doc.add_paragraph("Slide-ready architecture, frontend, backend, data-model and AI-control figures are supplied separately as 1920×1080 PNG, SVG and editable draw.io files.", style="Source Note")
    render_report(doc)
    doc.core_properties.title = "TreasuryPilot — Product & Technical Report"
    doc.core_properties.subject = "English report following the nine-part product and technical outline"
    doc.core_properties.author = "TreasuryPilot"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
