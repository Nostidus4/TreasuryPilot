#!/usr/bin/env python3
"""Rebuild the Excel progress tracker from the GitHub Project (GitHub is the source of truth).

Usage (repo root):
    uv run --with openpyxl python planning/tracker/build_tracker.py
Output: planning/tracker/TreasuryPilot-Tracker.xlsx
Requires: gh CLI logged in with the `project` scope.
"""
import json, os, subprocess, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
MAP = json.load(open(os.path.join(HERE, "..", "tools", "github_map.json")))
PROJECT_ID = MAP["project"]["id"]
OUT = os.path.join(HERE, "TreasuryPilot-Tracker.xlsx")
STATUSES = ["Draft", "Needs Discussion", "Ready for Review", "Accepted", "Ready for Assignment",
            "Assigned", "In Progress", "Review", "Done"]

Q = """query($p:ID!,$after:String){node(id:$p){...on ProjectV2{items(first:100,after:$after){
 pageInfo{hasNextPage endCursor}
 nodes{ fieldValues(first:20){nodes{
   ...on ProjectV2ItemFieldSingleSelectValue{name field{...on ProjectV2FieldCommon{name}}}
   ...on ProjectV2ItemFieldTextValue{text field{...on ProjectV2FieldCommon{name}}}}}
  content{...on Issue{number title url state assignees(first:5){nodes{login}} parent{number title}}}}}}}}"""

def fetch():
    rows, after = [], None
    while True:
        args = ["gh", "api", "graphql", "-f", f"query={Q}", "-f", f"p={PROJECT_ID}"]
        if after: args += ["-f", f"after={after}"]
        data = json.loads(subprocess.run(args, capture_output=True, text=True, check=True).stdout)
        items = data["data"]["node"]["items"]
        for n in items["nodes"]:
            c = n["content"] or {}
            f = {v["field"]["name"]: v.get("name", v.get("text")) for v in n["fieldValues"]["nodes"] if v and v.get("field")}
            rows.append({"id": f.get("Ticket ID", ""), "gh": c.get("number"), "title": c.get("title", ""), "url": c.get("url", ""),
                         "state": c.get("state", ""), "assignees": ", ".join(a["login"] for a in (c.get("assignees") or {}).get("nodes", [])),
                         "epic": (c.get("parent") or {}).get("title", ""), "module": f.get("Module", ""), "priority": f.get("Priority", ""),
                         "phase": f.get("Phase", ""), "size": f.get("Size", ""), "cut": f.get("Cut", ""), "area": f.get("Area", ""),
                         "status": "Done" if c.get("state") == "CLOSED" else f.get("Status", ""), "blocked_by": f.get("Blocked by", "") or ""})
        if not items["pageInfo"]["hasNextPage"]: break
        after = items["pageInfo"]["endCursor"]
    return rows

def build(rows):
    tickets = sorted([r for r in rows if not r["id"].startswith("EPIC-")], key=lambda r: r["id"])
    F = Font(name="Arial", size=10); FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    HEAD = PatternFill("solid", fgColor="1F3864"); YEL = PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()

    # ── Issues ──
    ws = wb.active; ws.title = "Issues"
    cols = ["Ticket ID", "GH #", "Title", "Module", "Epic", "Priority", "Phase", "Size", "Est. days", "Cut", "Area",
            "Status", "Assignees", "State", "Blocked by", "Open blockers", "Link"]
    ws.append(cols)
    for i, r in enumerate(tickets, start=2):
        ws.append([r["id"], r["gh"], r["title"], r["module"], r["epic"], r["priority"], r["phase"], r["size"],
                   f'=IF(H{i}="S",0.5,IF(H{i}="M",1,IF(H{i}="L",2,0)))', r["cut"], r["area"], r["status"],
                   r["assignees"], r["state"], r["blocked_by"], None, "open"])
        ws.cell(i, 17).hyperlink = r["url"]; ws.cell(i, 17).style = "Hyperlink"
    n = len(tickets) + 1
    # Open blockers: count of blockers whose status is not Done (looked up in this sheet)
    for i in range(2, n + 1):
        ws.cell(i, 16).value = (f'=IF(O{i}="",0,SUMPRODUCT(--ISNUMBER(SEARCH(", "&$A$2:$A${n}&",",", "&O{i}&",")),'
                                f'--($L$2:$L${n}<>"Done")))')
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"L2:L{n}")
    ws.conditional_formatting.add(f"L2:L{n}", CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(f"L2:L{n}", CellIsRule(operator="equal", formula=['"In Progress"'], fill=PatternFill("solid", fgColor="FFEB9C")))
    ws.conditional_formatting.add(f"P2:P{n}", FormulaRule(formula=[f'AND(P2>0,L2<>"Done")'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:Q{n}"

    # ── Module Progress ──
    mp = wb.create_sheet("Module Progress")
    modules = sorted({r["module"] for r in tickets if r["module"]})
    head = ["Module", "Tickets", "Done", "In Progress", "Review", "Not started", "% done (count)",
            "Est. days", "Est. days done", "% done (effort)", "Cut A tickets", "Cut A done", "% Cut A done"]
    mp.append(head)
    for j, mod in enumerate(modules, start=2):
        rng = f"Issues!$D$2:$D${n}"
        mp.append([mod,
                   f'=COUNTIFS({rng},A{j})',
                   f'=COUNTIFS({rng},A{j},Issues!$L$2:$L${n},"Done")',
                   f'=COUNTIFS({rng},A{j},Issues!$L$2:$L${n},"In Progress")',
                   f'=COUNTIFS({rng},A{j},Issues!$L$2:$L${n},"Review")',
                   f'=B{j}-C{j}-D{j}-E{j}',
                   f'=IF(B{j}=0,0,C{j}/B{j})',
                   f'=SUMIFS(Issues!$I$2:$I${n},{rng},A{j})',
                   f'=SUMIFS(Issues!$I$2:$I${n},{rng},A{j},Issues!$L$2:$L${n},"Done")',
                   f'=IF(H{j}=0,0,I{j}/H{j})',
                   f'=COUNTIFS({rng},A{j},Issues!$J$2:$J${n},"A")',
                   f'=COUNTIFS({rng},A{j},Issues!$J$2:$J${n},"A",Issues!$L$2:$L${n},"Done")',
                   f'=IF(K{j}=0,0,L{j}/K{j})'])
    t = len(modules) + 2
    mp.append(["TOTAL"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{t-1})" if c not in (7, 10, 13) else None for c in range(2, 14)])
    mp[f"G{t}"] = f"=IF(B{t}=0,0,C{t}/B{t})"; mp[f"J{t}"] = f"=IF(H{t}=0,0,I{t}/H{t})"; mp[f"M{t}"] = f"=IF(K{t}=0,0,L{t}/K{t})"
    for row in mp.iter_rows(min_row=2, max_row=t):
        for c in (row[6], row[9], row[12]): c.number_format = "0.0%"
    for c in mp[t]: c.font = Font(name="Arial", size=10, bold=True)

    # ── Status Summary ──
    ss = wb.create_sheet("Status Summary")
    ss.append(["Status", "All tickets", "Cut A", "Cut B", "M1", "Later"])
    for j, s in enumerate(STATUSES, start=2):
        L = f"Issues!$L$2:$L${n}"
        ss.append([s, f'=COUNTIFS({L},A{j})', f'=COUNTIFS({L},A{j},Issues!$J$2:$J${n},"A")',
                   f'=COUNTIFS({L},A{j},Issues!$J$2:$J${n},"B")', f'=COUNTIFS({L},A{j},Issues!$G$2:$G${n},"M1")',
                   f'=COUNTIFS({L},A{j},Issues!$G$2:$G${n},"Later")'])
    e = len(STATUSES) + 2
    ss.append(["TOTAL"] + [f"=SUM({c}2:{c}{e-1})" for c in "BCDEF"])
    ss.append([]); ss.append(["Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), "from GitHub Project", MAP["project"].get("url", "")])

    # ── Dependencies ──
    dp = wb.create_sheet("Dependencies")
    dp.append(["Ticket", "Blocked by", "Ticket status", "Blocker status", "Blocking now?"])
    k = 2
    for r in tickets:
        for b in [x.strip() for x in r["blocked_by"].split(",") if x.strip()]:
            dp.append([r["id"], b,
                       f'=IFERROR(INDEX(Issues!$L$2:$L${n},MATCH(A{k},Issues!$A$2:$A${n},0)),"")',
                       f'=IFERROR(INDEX(Issues!$L$2:$L${n},MATCH(B{k},Issues!$A$2:$A${n},0)),"")',
                       f'=IF(AND(C{k}<>"Done",D{k}<>"Done"),"Yes","No")'])
            k += 1
    dp.freeze_panes = "A2"; dp.auto_filter.ref = f"A1:E{k-1}"

    # ── How to use ──
    hw = wb.create_sheet("How to use")
    for line in ["This workbook is a REPORT. GitHub (Project 'TreasuryPilot · Delivery') is the source of truth.",
                 "To refresh: uv run --with openpyxl python planning/tracker/build_tracker.py",
                 "Do not edit statuses here; change them in the GitHub Project, then rebuild.",
                 "Status (column L, Issues) has a dropdown only for quick what-if analysis; changes are not synced back.",
                 "Est. days: S = 0.5, M = 1, L = 2 (planning estimates; not re-estimated by developers).",
                 "Open blockers = number of tickets listed in 'Blocked by' whose status is not Done.",
                 "Cut: A = required for the submission demo, B = stretch, After submission / Later = planned after 20/10."]:
        hw.append([line])

    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if c.hyperlink:
                    c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
                elif c.font is not None and c.font.bold:
                    c.font = Font(name="Arial", size=10, bold=True)
                else:
                    c.font = F
        if sh.title != "How to use":
            for c in sh[1]: c.font = FB; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in sh.columns:
            letter = get_column_letter(col[0].column)
            width = max(len(str(c.value)) if c.value and not str(c.value).startswith("=") else 8 for c in col)
            sh.column_dimensions[letter].width = min(max(10, width + 2), 70)
    wb["How to use"].column_dimensions["A"].width = 110
    wb.save(OUT)
    print(f"wrote {OUT} ({len(tickets)} tickets, {k-2} dependency rows)")

if __name__ == "__main__":
    build(fetch())
